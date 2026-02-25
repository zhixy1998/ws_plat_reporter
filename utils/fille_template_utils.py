from math import sin
import os


import re
import base64
import importlib
import urllib.parse
from io import BytesIO
from decimal import Decimal
from typing import List, Dict
from dotenv import load_dotenv

import requests
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxtpl import InlineImage, DocxTemplate
from app.template_doc.model.template_doc import TemplateDoc
from config import config

from app.db import execute_safe_sql
from exceptions.error import ErrorCode
from constants.template_constant import MeasureType, ExperimentType, TemplateType
from exceptions.command_exception import APIException
from app.template_doc.model.template_bookmark import TemplateBookmark
from constants.template_constant import GetTypeConstant, FillTypeConstant
from utils.table_data_processor import TableDataProcessor, GetHeaderAndData
from utils.classic_fille_template_utils_simple import ClassicFillDataUtils
from utils.template_utils import remove_marker_rows

load_dotenv(override=True)

class FillDataUtils():
    def __init__(self, name,  fake=False, new_title_mapping=None, template=None, structure_path=config.structure_module):
        self.name = name
        self.sample_id = None
        private_mod = None
        self.common_val = None   # 代表通用数据接口返回的数据
        self.common_key = None   # 代表通用数据的描述
        self.private_key = None  # 代表私有数据的描述
        self.python_jinja_mapping = {}
        self.table_contexts = {}
        self.contexts = {}
        self.end_point_header_and_data = None
        self.fake = fake
        self.table_data_processor = TableDataProcessor(name, structure_path=structure_path)
        self.init_data(structure_path)
        self.new_title_mapping = new_title_mapping
        self.single_sample = False
        self.template_type = None
        self.template = template


    def convert_tel_val_data_to_rows(self, tel_val_data_list):
        """
        将tel_val_data_list转换为行格式
        输入: {'水样体积': ['6', '7'], '空白试验吸光度': ['4', '6'], ...}
        输出: [{'水样体积': '6', '空白试验吸光度': '4', ...}, {'水样体积': '7', '空白试验吸光度': '6', ...}]
        """
        if not tel_val_data_list:
            return []
        
        # 找到最长的数据列表长度
        max_length = max(len(values) for values in tel_val_data_list.values()) if tel_val_data_list else 0
        
        # 构建行数据
        rows = []
        for i in range(max_length):
            row = {}
            for key, values in tel_val_data_list.items():
                # 如果下标溢出，使用空字符串
                row[key] = values[i] if i < len(values) else ''
            rows.append(row)
        
        return rows

    def init_data(self, structure_path):
        common_data_module_path = f"{structure_path}.common_data"  # 比如 name = "my_template"
        common_mod = importlib.import_module(common_data_module_path)
        self.common_title_mapping = common_mod.title_mapping
        common_key = common_mod.data
        self.common_key = common_key
        module_name = f"{structure_path}.{self.name}"  # 比如 name = "my_template"
        try:
            private_mod = importlib.import_module(module_name)
        except Exception as e:
            raise APIException(error=ErrorCode.TEMPLATE_NOT_FOUND)
        self.private_mod = private_mod
        # 不是表格数据
        self.private_key = private_mod.python_template_var
        

    async def fill_data(self, doc_stream, task_id, sample_id=None, template_id=None, analysis_id=None, image_base64=None):
        common_val = self.get_common_data_by_endpoint(task_id, analysis_id=analysis_id)

        # get_table_head_data_from_endpoint expects (task_id, analysis_id, ...)
        end_point_header_and_data = self.get_table_head_data_from_endpoint(task_id, analysis_id)
        self.end_point_header_and_data = end_point_header_and_data
        # 先加载通用数据

        private_key_text_fields = self.private_mod.python_template_var    
        # 表格数据
        # private_data_table_fields = private_mod.table_fields
        # self.table_fields = private_mod.table_fields

        get_data_from_common_data_jinja = [] # 这个数据从通用里面取 然后jinja填充
        get_data_from_common_data_bookmark = [] # 这个数据从通用里面取 然后书签填充 
        get_data_from_endpoint_jinja = [] # 这个数据从接口取 然后jinja填充
        get_data_from_endpoint_bookmark = [] # 这个数据从接口取 然后书签填充
        get_data_from_database_jinja = [] # 这个数据从数据库取 然后jinja填充
        get_data_from_database_bookmark = [] # 这个数据从数据库取 然后书签填充

        for k, item in private_key_text_fields.items():
            if item.get('get_type') is None or int(item.get('get_type'))  == GetTypeConstant.common.value:
                if int(item.get('fill_type')) == FillTypeConstant.jinja.value:
                    get_data_from_common_data_jinja.append({k:item})
                if int(item.get('fill_type')) == FillTypeConstant.bookmark.value:
                    get_data_from_common_data_bookmark.append({k:item})
                
            elif int(item.get('get_type')) == GetTypeConstant.endpoint.value:
                if int(item.get('fill_type')) == FillTypeConstant.jinja.value:
                    get_data_from_endpoint_jinja.append({k:item})
                if int(item.get('fill_type')) == FillTypeConstant.bookmark.value:
                    get_data_from_endpoint_bookmark.append({k:item})
            elif int(item.get('get_type')) == GetTypeConstant.database.value:
                if int(item.get('fill_type')) == FillTypeConstant.jinja.value:
                    get_data_from_database_jinja.append({k:item})
                if int(item.get('fill_type')) == FillTypeConstant.bookmark.value:
                    get_data_from_database_bookmark.append({k:item})
            
        if get_data_from_common_data_bookmark:
            pass
        if get_data_from_endpoint_jinja:
            pass
        if get_data_from_endpoint_bookmark:
            pass
        if get_data_from_database_jinja:
            pass
        if get_data_from_database_bookmark:
            pass
    
        template_type = self.template.type
        self.template_type = template_type
        self.experiment_type = self.template.experiment_type
        await self.handle_measure_type(task_id=task_id)
        if analysis_id and len(analysis_id.split(',')) == 1:
            self.single_sample = True

        # 原始记录报告 走的逻辑
        await self.fill_common_data_title(doc_stream)
        # 确保 tpl 在所有分支中都有定义，避免后续无条件使用时未声明报错
        tpl = None
        if self.template_type == TemplateType.report_template.value:
            await self.get_table_context(doc_stream, task_id)
            await self.fill_common_data(doc_stream=doc_stream)
        # 测量数据走的逻辑
        if self.template_type == TemplateType.original_records.value:
            classic_fill_data_utils = ClassicFillDataUtils(self.name, self.template.id, self.measure_type, single_sample=self.single_sample)
            tpl = DocxTemplate(doc_stream)
            contexts = await classic_fill_data_utils.guide_get_context(self.common_val, self.end_point_header_and_data, tpl=tpl, image_base64=image_base64)
            self.contexts.update(contexts)
        doc_stream = self.fill_template_by_jinja(doc_stream, tpl=tpl)
        doc_stream = self.insert_header_page_number(doc_stream)
        return doc_stream

    
    async def handle_measure_type(self, task_id=None):
        measure_type = None
        measure_type = await self.get_measure_type(task_id)
        if measure_type is None and self.experiment_type is not None:
            measure_type = ExperimentType.get_name_desc_by_value(self.experiment_type)
        self.measure_type = measure_type

    async def preview_view(self, doc_stream, task_id=None, remark=None, image_base64=None):

        template_type = self.template.type
        self.template_type = template_type
        self.experiment_type = self.template.experiment_type
        self.measure_type = ExperimentType.get_name_desc_by_value(self.experiment_type)
        if template_type == TemplateType.report_template.value:
            # 获取表格上下文与通用数据
            await self.get_table_context(doc_stream=doc_stream)
        if remark:
            await self.fill_remark(remark)
        await self.fill_common_data(doc_stream=doc_stream, title=True, v=True)

        if self.measure_type == MeasureType.spectral_scanning.decs: # 光谱扫描
            await self.get_spectral_scanning_table_context()
        elif self.measure_type == MeasureType.time_scanning.decs: # 时间扫描
            await self.get_time_scanning_table_context()
        elif self.measure_type == MeasureType.quantitative_determination.decs: # 定量测定
            await self.get_quantitative_determination_table_context()
        elif self.measure_type == MeasureType.photometric_measurement.decs: # 光度测量
            await self.get_photometric_measurement_table_context()

        # DEBUG: inspect incoming stream
        try:
            doc_stream.seek(0)
            incoming_size = len(doc_stream.getvalue())
        except Exception:
            incoming_size = None
        print(f"DEBUG preview_view: before fill_template_by_jinja, type={type(doc_stream)}, size={incoming_size}")

        try:
            doc_stream = self.fill_template_by_jinja(doc_stream=doc_stream)
        except Exception as e:
            print(f"ERROR preview_view: fill_template_by_jinja raised: {e}")
            raise

        # DEBUG: inspect returned stream
        try:
            doc_stream.seek(0)
            out_size = len(doc_stream.getvalue())
        except Exception:
            out_size = None
        print(f"DEBUG preview_view: after fill_template_by_jinja, type={type(doc_stream)}, size={out_size}")

        return doc_stream

    async def get_table_context(self, doc_stream, task_id=None, image_base64=None):
        # 原始记录报告
        if task_id:
            await self.handle_measure_type(task_id)
        else:
            if self.template:
                self.measure_type = ExperimentType.get_name_desc_by_value(self.template.experiment_type)
            if self.template_type is None:
                self.template_type = self.template.type

        # 定量和光度测量
        if self.measure_type == MeasureType.quantitative_determination.decs:
            self.HEADER_M_MAPPING_LIST = self.table_data_processor.handler_instance.HEADER_M_MAPPING_LIST  # 每个私有文件都要声明的
            self.TABLE_HEADER_MAP  = self.table_data_processor.handler_instance.TABLE_HEADER_MAP
            photometric_measurement_context = await self.get_quantitative_determination_photometric_measurement_from_endpoint(doc_stream)
            # origin_header_real_header_mapping, tpl_context = self.get_quantitative_determination_photometric_measurement_from_database(doc_stream, task_id)
            self.contexts.update(photometric_measurement_context)
        else:
            pass
    
    async def fill_common_data(self, doc_stream, title=False, v=True):
        # 填充
        if title:
            await self.fill_common_data_title(doc_stream)
        if v:
            await self.fill_common_data_v()

    async def fill_common_data_title(self, doc_stream):

        has_private_title_mapping = False
        if hasattr(self.table_data_processor.template_module, 'title_mapping'):
            has_private_title_mapping = True
            private_title_mapping = self.table_data_processor.template_module.title_mapping
        doc_stream.seek(0)
        doc = DocxTemplate(doc_stream)
        context_vars = doc.get_undeclared_template_variables()
        if self.new_title_mapping:
            for t in self.new_title_mapping:
                if t.endswith('_t'):
                    v = self.new_title_mapping.get(t).get('name', '')
                    if v:
                        self.contexts[t] = v
        else:
            for t in context_vars:
                if t.endswith('_t'):
                    if has_private_title_mapping:
                        v = private_title_mapping.get(t).get('name', '')
                        if not v:
                            v = self.common_title_mapping.get(v).get('name', '')
                        if v:
                            self.contexts[t] = v
                    else:
                        v = self.common_title_mapping.get(t).get('name', '')
                        if v:
                            self.contexts[t] = v
    
    async def fill_common_data_v(self):
        if self.template_type == TemplateType.report_template.value:
            if not self.fake:
                for private_k, v_dict in self.private_key.items():
                    if not private_k.endswith('_t'): # 不是标题代表数值
                        v = self.common_val.get(private_k, '')
                        if v:
                            self.contexts[private_k] = v
                        if v_dict.get('position'):
                            pass
                        else:
                        # 查看通用数据是否描述了这个变量的取值方法。
                            if private_k in self.common_key:
                                get_desc = self.common_key[private_k]
                                if get_desc.get('position'):
                                    await self.handle_position(private_k, get_desc.get('position'))
                                if get_desc.get('children'):
                                    await self.handle_children(private_k, get_desc.get('children'), connect_str=get_desc.get('connect_str', ''))
                            else:
                                if not self.contexts.get(private_k):
                                    self.contexts[private_k] = ''
            else:
                for private_k, v_dict in self.private_key.items():
                    if not private_k.endswith('_t') and not self.contexts.get(private_k): # 不是标题代表数值
                        self.contexts[private_k] = ' '

        
        

    # async def get_blank_table_context(self, template_name):
    #     origin_header_real_header_mapping = {} # {'模版本来的表头': '实际表头'}
    #     table_context = self.get_table_context()
    #     return tpl_context



    def get_common_data_by_endpoint(self, task_id, analysis_id=None):
        template_common_url = config.guide_template_common_var_url + '?taskId=' + urllib.parse.quote(str(task_id), safe='') + '&analysisId=' + urllib.parse.quote(str(analysis_id), safe='')
        # print(f"DEBUG: 请求URL: {template_common_url}")
        # print(f"DEBUG: 原始task_id: {repr(task_id)}")
        try:
            res = requests.get(template_common_url, timeout=5)
            if res.status_code == 200:
                try:
                    res_json = res.json()
                    if res_json.get('code') == 200:
                        common_val = res_json.get('data')
                    else:
                        common_val = {}
                except Exception as e:
                    print(e)
                    common_val = {}
            else:
                print("请求通用数据失败, url:", template_common_url)
                common_val = {}
        except requests.exceptions.Timeout:
            print(f"请求超时，跳过: {template_common_url}")
            common_val = {}
        except Exception as e:
            print(f"请求异常，跳过: {e}")
            common_val = {}
        
        self.common_val = common_val
        return common_val

    # 通用数据的jinja填充数据
    # async def common_fill_data_by_jinja(self, doc_stream, additional_context=None):
    #     """
    #     逻辑: 如果private里面有，但是没有描述position, 那么就从common里面取, 如果描述了position，再说
    #     """
    #     if additional_context is None:
    #         additional_context = {}

    #     # 确保流位置在开始
    #     if hasattr(doc_stream, 'seek'):
    #         doc_stream.seek(0)

    #     tpl = DocxTemplate(doc_stream)
    #     value_list = [] # 所有的值的列表
    #     for private_k, v_dict in self.private_key.items():
    #         # 如果私有文件描述了怎么取数据。
    #         if v_dict.get('position'):
    #             pass
    #         else:
    #         # 查看通用数据是否描述了这个变量的取值方法。
    #             if private_k in self.common_key:
    #                 get_desc = self.common_key[private_k]
    #                 if get_desc.get('position'):
    #                     await self.handle_position(private_k, get_desc.get('position'), value_list)
    #                 if get_desc.get('children'):
    #                     await self.handle_children(private_k, get_desc.get('children'), value_list, connect_str=get_desc.get('connect_str', ''))
    #             else:
    #                 self.python_jinja_mapping[private_k] = ''
    #     for k, v in self.python_jinja_mapping.items():
    #         # print(k, v)
    #         if v == '':
    #             self.python_jinja_mapping[k] = ' '
    #     intermediate_stream = self.fill_template_by_jinja(tpl, additional_context)
    #     return intermediate_stream

    # 从通用接口取当前表格需要的jinja数据
    async def get_common_val_by_private_position(self, position):
        val = self.common_val.get(position[0])
        return val

    # 处理position position代表这个数据获取的描述方法。
    async def handle_position(self, private_k, position):
        val = await self.get_common_val_by_private_position(position)
        if private_k == 'b':
            self.contexts[private_k] = val.split(',')[0] if val else ''
        elif private_k == 'a':
            try:
                self.contexts[private_k] = val.split(',')[1] if val else ''
            except Exception:
                self.contexts[private_k] = val if val else ''
        # if '图片' in private_k and isinstance(val, str) and self.is_base64(val):
        #     try:
        #         self.python_jinja_mapping[private_k] = await self.decode_base64_image(bookmark_name=private_k, tpl=tpl, val=val)
        #     except Exception as e:
        #         print(f"{private_k} 解码失败: {e}")
        # else:
        else:
            self.contexts[private_k] = val if val else ''
    
    # 处理children children代表这个数据获取的描述方法。
    async def handle_children(self, private_k, children, connect_str=''):
        final_val_list = []
        for child in children:
            get_desc = self.common_key.get(child)
            if get_desc and get_desc.get('position'):
                val = await self.get_common_val_by_private_position(get_desc.get('position'))
                if val:
                    final_val_list.append(val)
        if final_val_list:
            self.contexts[private_k] = connect_str.join(final_val_list)
        


    # 通用的通过书签填充
    def common_fill_data_by_bookmark(private_key, common_key, common_val, doc_stream):
        pass

    # 通过接口获取数据然后jinja填充
    def endpoint_fill_data_by_jinja(self):
        pass

    # 通过接口获取数据然后书签填充
    def endpoint_fill_data_by_bookmark(self):
        pass

    # 通过数据库获取数据然后jinja填充
    def database_fill_data_by_jinja(self):
        pass

    # 通过数据库获取数据然后书签填充
    def database_fill_data_by_bookmark(self):
        pass

    def fill_template_by_jinja(self, doc_stream: BytesIO, tpl=None) -> BytesIO:
        """
        合并 context 并用 docxtpl 渲染模板。
        如果传入了 DocxTemplate 实例 `tpl`，优先使用它（保证 subdoc 归属一致）。
        否则根据 `doc_stream` 新建一个 `DocxTemplate`。
        返回渲染后的 BytesIO（出错时尽量返回原始流的副本以避免 None）。
        """
        try:
            print(f"DEBUG fill_template_by_jinja: enter, doc_stream type={type(doc_stream)}, tpl_type={type(tpl)}")
            from docxtpl import DocxTemplate as _DocxTemplate

            used_tpl = None
            if tpl is not None:
                # 如果传入的是 DocxTemplate 实例，直接使用
                if isinstance(tpl, _DocxTemplate):
                    used_tpl = tpl
                else:
                    # 尝试将其构造为 DocxTemplate（tpl 可能是 path/stream）
                    try:
                        used_tpl = _DocxTemplate(tpl)
                    except Exception:
                        used_tpl = None

            if used_tpl is None:
                try:
                    doc_stream.seek(0)
                except Exception:
                    pass
                used_tpl = _DocxTemplate(doc_stream)
            undeclared_key = [i for i in used_tpl.get_undeclared_template_variables() if i.endswith('_t')]
            declared_key = [i for i in self.contexts.keys() if i.endswith('_t')]

            # 找出模板中有但 contexts 中没有数据的 _t 变量，打上删除标记
            _DEL_MARKER = '__DEL_ROW__'
            extra_keys = set(undeclared_key) - set(declared_key)
            if extra_keys:
                print(f"DEBUG fill_template_by_jinja: 多余的_t变量(将标记删除行): {extra_keys}")
                for ek in extra_keys:
                    self.contexts[ek] = _DEL_MARKER

            # 合并context 并保存
            used_tpl.render(context=self.contexts)
            output_stream = BytesIO()
            used_tpl.save(output_stream)
            output_stream.seek(0)

            # 删除含有删除标记的表格行（只删标记所在单元格的行，不影响左侧分页）
            output_stream = remove_marker_rows(output_stream, _DEL_MARKER)

            # 移除表格中的空行
            output_stream = self.remove_empty_table_rows(output_stream)
            try:
                print(f"DEBUG fill_template_by_jinja: success, size={len(output_stream.getvalue())}")
            except Exception:
                pass
            return output_stream

        except Exception as e:
            print(f"ERROR fill_template_by_jinja: {e}")
            # 尝试返回原始流的副本，避免返回 None
            try:
                try:
                    doc_stream.seek(0)
                except Exception:
                    pass
                if hasattr(doc_stream, 'getvalue'):
                    return BytesIO(doc_stream.getvalue())
                elif hasattr(doc_stream, 'read'):
                    try:
                        doc_stream.seek(0)
                    except Exception:
                        pass
                    data = doc_stream.read()
                    if data is not None:
                        return BytesIO(data)
            except Exception:
                pass
            return doc_stream
    

    def remove_empty_table_rows(self, doc_stream: BytesIO) -> BytesIO:
        """移除表格中的空行"""
        try:
            from docx import Document
            doc_stream.seek(0)
            doc = Document(doc_stream)
            
            for table in doc.tables:
                rows_to_remove = []
                for i, row in enumerate(table.rows):
                    # 检查行是否为空（所有单元格都为空或只包含空白字符）
                    is_empty = True
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            is_empty = False
                            break
                    
                    if is_empty:
                        rows_to_remove.append(i)
                
                # 从后往前删除空行，避免索引问题
                for row_idx in reversed(rows_to_remove):
                    if row_idx < len(table.rows):
                        table._element.remove(table.rows[row_idx]._element)
            
            # 保存处理后的文档
            output_stream = BytesIO()
            doc.save(output_stream)
            output_stream.seek(0)
            return output_stream
            
        except Exception as e:
            print(f"Error removing empty table rows: {e}")
            doc_stream.seek(0)
            return doc_stream
    
    # def get_common_table_data(self, table_name: str, task_id: str) -> dict:
    #     template_common_url = os.getenv('TEMPLATE_COMMON_TABLE_URL') + '?taskId=' + urllib.parse.quote(str(task_id), safe='') + '&reportName=' + table_name + "&sampleId=" + urllib.parse.quote(str(self.sample_id), safe='')
    #     print(f"DEBUG: 请求URL: {template_common_url}")
    #     print(f"DEBUG: 原始task_id: {repr(task_id)}")
    #     try:
    #         res = requests.get(template_common_url, timeout=5)
    #         if res.status_code == 200:
    #             common_table_val = res.json().get('data')
    #         else:
    #             print("请求通用表格数据失败, url:", template_common_url)
    #             common_table_val = {}
    #     except requests.exceptions.Timeout:
    #         print(f"请求超时，跳过: {template_common_url}")
    #         common_table_val = {}
    #     except Exception as e:
    #         print(f"请求异常，跳过: {e}")
    #         common_table_val = {}
    #     return common_table_val
    
    def handle_endpoint_data_to_table_row(self, table_fields: list, common_table_val) -> list:
        pass

    def is_base64(self, s: str) -> bool:
        if not isinstance(s, str):
            return False
        try:
            # base64 字符长度必须是4的倍数
            if len(s) % 4 != 0:
                return False
            # 符合 base64 编码字符集
            if not re.match(r'^[A-Za-z0-9+/=\s]+$', s):
                return False
            # 能否成功 decode
            base64.b64decode(s, validate=True)
            return True
        except Exception:
            return False

    async def fill_remark(self, remark_item):
        self.contexts['remark'] = remark_item.value
    
    async def decode_base64_image(self, bookmark_name, tpl, val):
        try:
            image_data = base64.b64decode(val)
            image_stream = BytesIO(image_data)
            img_bookmark_instance = await TemplateBookmark.get_or_none(template_doc_id=self.template_id, name=bookmark_name, is_del=0)
            if img_bookmark_instance and img_bookmark_instance.img_size:
                return InlineImage(tpl, image_stream, width=Cm(img_bookmark_instance.img_size))  # 根据需要调整宽度
            else:
                return InlineImage(tpl, image_stream, width=Cm(14))  # 根据需要调整宽度
        except Exception as e:
            print("图片解码失败", e)
            return None
    
    def filter_table_header(self, private_header, table_header_data_mapping):
        private_header = self.handle_table_header(private_header)
        new_table_header_data_mapping = {}
        for header in private_header:
            if header in table_header_data_mapping:
                new_table_header_data_mapping[header] = table_header_data_mapping[header]
        return new_table_header_data_mapping
    
    def handle_table_header(self, headers):
        """
        数据库存储的表头可能是嵌套的，需要处理
        两层嵌套，就用 '-' 把同一列连接起来。
        """
        if len(headers) == 1:
            return headers[0]
        else:
            return ["-".join(map(str, group)) for group in zip(*headers)]
    

    def insert_header_page_number(self, doc_stream: BytesIO, alignment: int = 3) -> BytesIO:
        """
        在每个 section 的页脚追加页码（不覆盖已有 header/footer 内容），并确保每节都能显示。

        Args:
            doc_stream: Word文档流
            alignment: 对齐方式 (1:左对齐, 2:居中, 3:右对齐)

        Returns:
            BytesIO: 修改后的文档流
        """
        try:
            print(f"DEBUG insert_header_page_number: enter, doc_stream type={type(doc_stream)}, alignment={alignment}")
            doc_stream.seek(0)
            doc = Document(doc_stream)

            # 转换对齐方式：1:左, 2:中, 3:右 -> 0:左, 1:中, 2:右
            alignment_map = {1: 0, 2: 1, 3: 2}
            para_align = alignment_map.get(alignment, 2)

            def add_field(p, field_code):
                run = p.add_run()
                fldChar_begin = OxmlElement('w:fldChar')
                fldChar_begin.set(qn('w:fldCharType'), 'begin')
                instrText = OxmlElement('w:instrText')
                instrText.set(qn('xml:space'), 'preserve')
                instrText.text = field_code
                fldChar_separate = OxmlElement('w:fldChar')
                fldChar_separate.set(qn('w:fldCharType'), 'separate')
                fldChar_end = OxmlElement('w:fldChar')
                fldChar_end.set(qn('w:fldCharType'), 'end')
                run._r.append(fldChar_begin)
                run._r.append(instrText)
                run._r.append(fldChar_separate)
                run._r.append(fldChar_end)

            # 在每个 section 的 footer 中追加页码段落，且确保节设置不阻止显示
            for section in doc.sections:
                # 关闭不同首页/奇偶页导致的显示差异
                try:
                    section.different_first_page_header_footer = False
                except Exception:
                    pass

                # 断开与上一节的 header/footer 链接（若有）
                try:
                    # header/footer 对象可能没有 is_linked_to_previous 属性，安全设置
                    if hasattr(section.header, 'is_linked_to_previous'):
                        section.header.is_linked_to_previous = False
                    if hasattr(section.footer, 'is_linked_to_previous'):
                        section.footer.is_linked_to_previous = False
                except Exception:
                    pass

                footer = section.footer
                # 在 footer 末尾追加一段，不删除已有内容
                para = footer.add_paragraph()
                para.alignment = para_align
                para.add_run("第 ")
                add_field(para, "PAGE")
                para.add_run(" 页，共 ")
                add_field(para, "NUMPAGES")
                para.add_run(" 页")

            # 保存回字节流
            output_stream = BytesIO()
            doc.save(output_stream)
            output_stream.seek(0)
            try:
                print(f"DEBUG insert_header_page_number: success, size={len(output_stream.getvalue())}")
            except Exception:
                pass
            return output_stream
        except Exception as e:
            print(f"ERROR insert_header_page_number: {e}")
            try:
                doc_stream.seek(0)
            except Exception:
                pass
            return doc_stream
    
    def remove_header_page_number(self, doc_stream: BytesIO) -> BytesIO:
        """
        从所有节的 footer 中移除包含 PAGE/NUMPAGES 域的段落（对追加式插入的页码有效）。

        Args:
            doc_stream: Word文档流

        Returns:
            BytesIO: 修改后的文档流
        """
        doc_stream.seek(0)
        doc = Document(doc_stream)

        def paragraph_contains_page_fields(paragraph):
            # 检查段落中的任何 run 是否包含 fld instrText 包含 PAGE 或 NUMPAGES
            for run in paragraph.runs:
                for child in run._r:
                    if child.tag.endswith('}instrText'):
                        text = child.text or ''
                        if 'PAGE' in text or 'NUMPAGES' in text:
                            return True
            # 也检查文本层面（兼容旧实现）
            txt = paragraph.text or ''
            if '第' in txt and '页' in txt and '共' in txt:
                return True
            return False

        for section in doc.sections:
            footer = section.footer
            # 收集要删除的段落元素
            paras_to_remove = []
            for para in footer.paragraphs:
                if paragraph_contains_page_fields(para):
                    paras_to_remove.append(para)

            for para in paras_to_remove:
                try:
                    p = para._element
                    p.getparent().remove(p)
                except Exception:
                    pass

        # 保存回字节流
        output_stream = BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)
        return output_stream
    
    
    async def get_all_child_bookmarks(self, parent_id):
        """
        递归获取所有子标签，组成一维列表
        
        Args:
            parent_id: 父标签的ID
            
        Returns:
            list: 所有子标签的一维列表
        """
        all_children = []
        
        # 获取直接子标签
        direct_children = await TemplateBookmark.filter(
            parent_id=parent_id, 
            template_doc_id=self.template_id, 
            is_del=0
        ).all()
        
        for child in direct_children:
            # 添加当前子标签到列表
            all_children.append(child)
            
            # 递归获取这个子标签的所有子标签
            grandchildren = await self.get_all_child_bookmarks(child.id)
            all_children.extend(grandchildren)
        
        return all_children
    
    def fill_template_jinja_custom_info(self,doc_stream, custom_info=None, additional_context=None):
        """
        doc_stream    模板文件流
        custom_info   自定义标签信息
        """
        module_name = "template_structure." + self.name
        try:
            private_mod = importlib.import_module(module_name)
        except Exception as e:
            private_mod = None
        # if private_mod and hasattr(private_mod, 'custom_bookmark'):
        #     custom_bookmark=private_mod.custom_bookmark
        # else:
        #     custom_bookmark = {'note': '备注', 'project_name': '项目名称'}

        try:
            from docxtpl import DocxTemplate
            from io import BytesIO

            # 创建DocxTemplate对象
            doc_stream.seek(0)
            tpl = DocxTemplate(doc_stream)

            # 获取模板中所有未声明的变量
            undeclared_vars = tpl.get_undeclared_template_variables()

            blank_context = {}
            for var in undeclared_vars:
                blank_context[var] = ""
            # if custom_info:
            #     for i in undeclared_vars:
            #         if i in custom_bookmark['project_name']:
            #             blank_context[i] = custom_info.get('project_name')
            #         if i in custom_bookmark['note']:
            #             blank_context[i] = custom_info.get('note')

            # 为所有未声明的变量创建空字符串映射

            # 如果没有未声明的变量，直接返回原文档
            if not blank_context:
                doc_stream.seek(0)
                return doc_stream
            if additional_context:
                blank_context.update(additional_context)
            # 使用空字符串上下文渲染模板
            tpl.render(blank_context)

            # 保存到新的BytesIO流
            output_stream = BytesIO()
            tpl.save(output_stream)
            output_stream.seek(0)

            return output_stream

        except Exception as e:
            # 出错时返回原文档
            doc_stream.seek(0)
            return doc_stream
    
    
    @staticmethod
    async def get_measure_type(task_id):
        if not task_id:
            return None
        sql = "select measure_type from t_measure_task where id = %s"
        sql_result = await execute_safe_sql(sql, [task_id], return_dict=True)
        if sql_result:
            return sql_result[0].get('measure_type')
        return None
    
    async def get_quantitative_determination_photometric_measurement_from_endpoint(self, doc_stream, origin_header_real_header_mapping={}, tpl_context={}):
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        
        temp_context = {}
        self.known_sample_jinja_v = self.table_data_processor.get_known_sample_jinja_v()
        
        # 方案1：直接在外层单元格中嵌套表格，不使用subdoc
        doc_stream.seek(0)
        outer_doc = Document(doc_stream)
        
        # 处理已知样品
        if self.table_data_processor.handler_instance:
            stand = getattr(self.table_data_processor.handler_instance, 'STAND')
            if stand:
                head = []
                row = []
                if not self.fake:
                    for s in stand:
                        real_header= self.get_stand_real_header_from_endpoint_data(s)
                        head.append(real_header)
                        data = self.get_stand_data_list(s)
                        row.append(data)
                    for index, row_data in enumerate(row):
                        row_data.insert(0, head[index])
                else:
                    for index, s in enumerate(stand):
                        real_header= f'已知样-{index}'
                        head.append(real_header)
                        data = ['样品'+str(i+1) for i in range(6)]
                        row.append(data)
                    for index, row_data in enumerate(row):
                        row_data.insert(0, head[index])
                
                # 处理已知样品表格
                rows = len(row)
                cols = max(len(r) for r in row) if rows else 0
                
                if rows > 0 and cols > 0:
                    # 计算每行的高度
                    from docx.shared import Inches
                    row_height = int(Inches(0.25))  # 每行高度
                    
                    # 模板默认支持2行，计算是否需要额外的高度
                    default_rows = 2
                    extra_height = 0
                    height_multiplier = 1.5  # 增加高度的倍数系数
                    if rows > default_rows:
                        # 需要额外增加的高度（乘以倍数系数）
                        extra_height = int((rows - default_rows) * row_height * height_multiplier)
                        # print(f"DEBUG: 嵌入表格行数({rows}) > 默认行数({default_rows})，需增加高度: {extra_height} EMU (倍数系数: {height_multiplier})")
                    
                    # 在外层表格的单元格中查找并嵌入表格
                    found_and_embedded = False
                    for table in outer_doc.tables:
                        for row_idx, row_obj in enumerate(table.rows):
                            for cell_idx, cell in enumerate(row_obj.cells):
                                # 查找包含placeholder的单元格
                                if self.known_sample_jinja_v in cell.text:
                                    # 第一步：对每一行设置高度（包括增加量的分配）
                                    # 计算每行应该增加的高度
                                    height_per_row = row_height
                                    if extra_height > 0:
                                        # 将增加量平均分配到每一行
                                        height_per_row = row_height + (extra_height // rows)
                                    
                                    # 设置从row_idx开始的rows行的高度
                                    for offset in range(rows):
                                        if row_idx + offset < len(table.rows):
                                            table.rows[row_idx + offset].height = height_per_row
                                            # print(f"DEBUG: 设置外层表格第 {row_idx + offset + 1} 行高度为: {height_per_row} EMU")
                                    
                                    # 清空该单元格的所有内容
                                    cell.text = ""
                                    for para in cell.paragraphs:
                                        p = para._element
                                        p.getparent().remove(p)
                                    
                                    # 第二步：在单元格中添加新表格
                                    inner_table = cell.add_table(rows=rows, cols=cols)
                                    inner_table.style = "Table Grid"
                                    
                                    # 设置表格样式
                                    self._remove_outer_borders_only(inner_table)
                                    self._set_cell_margins(inner_table, margin=0)
                                    self._set_table_full_width(inner_table)
                                    self._set_first_column_wider(inner_table, first_col_pct=12)
                                    
                                    # 填充表格数据
                                    for r, row_data in enumerate(row):
                                        inner_table.rows[r].height = row_height
                                        for c in range(cols):
                                            cell_obj = inner_table.cell(r, c)
                                            cell_obj.text = "" if c >= len(row_data) or row_data[c] is None else str(row_data[c])
                                            
                                            # 设置单元格垂直对齐为顶部
                                            tc = cell_obj._tc
                                            tcPr = tc.get_or_add_tcPr()
                                            tcVAlign = OxmlElement('w:vAlign')
                                            tcVAlign.set(qn('w:val'), 'top')
                                            for existing in tcPr.findall(qn('w:vAlign')):
                                                tcPr.remove(existing)
                                            tcPr.append(tcVAlign)
                                            
                                            # 清除单元格间距
                                            for para in cell_obj.paragraphs:
                                                para.paragraph_format.space_before = 0
                                                para.paragraph_format.space_after = 0
                                    
                                    # 清除内嵌表格的段落间距
                                    for para in cell.paragraphs:
                                        para.paragraph_format.space_before = 0
                                        para.paragraph_format.space_after = 0
                                    
                                    # print(f"DEBUG: 在外层单元格[行{row_idx}, 列{cell_idx}]中嵌入表格 (行数: {rows}, 列数: {cols})")
                                    # print(f"DEBUG: 每行高度: {height_per_row} EMU (基础: {row_height}, 增加: {extra_height // rows if extra_height > 0 else 0})")
                                    found_and_embedded = True
                                    break
                            if found_and_embedded:
                                break
                        if found_and_embedded:
                            break
    
        # 无论是否处理了已知样品，都需要保存文档回流（因为Document已经加载了doc_stream）
        doc_stream.seek(0)
        doc_stream.truncate(0)  # 清空流内容
        outer_doc.save(doc_stream)
        doc_stream.seek(0)
        # print(f"DEBUG: 已保存文档到doc_stream")

        # 处理未知样品
        if self.table_data_processor.handler_instance:
            # total_handler_list = table_data_processor.list_available_handlers()
            # for handler in total_handler_list:
            #     handler()
            # standard = table_data_processor.handler_instance.standard
            # if not standard:
            #     standard = []
            if hasattr(self.table_data_processor.handler_instance, 'handle_unknown_sample_table'):
                private_context = self.table_data_processor.handler_instance.handle_unknown_sample_table(fake=self.fake, end_point_header_and_data=self.end_point_header_and_data)
                self.contexts.update(private_context)
            else:
                for index, origin_header_map in enumerate(self.HEADER_M_MAPPING_LIST):
                    table_name = 'table_' + str(index+1)
                    table_origin_header_data_mapping = {}  # 每个表格单独的数据映射
                    for jinja_header_value, m in origin_header_map.items():
                        # 表头映射
                        table_title_k = jinja_header_value + '_h'
                        if table_title_k not in self.TABLE_HEADER_MAP:
                            print(f'数据字段{jinja_header_value}没有在TABLE_HEADER_MAP中定义，跳过该字段。 应该定义{table_title_k}')
                            continue
                        if not self.fake:
                            real_header, identity = self.get_unknown_sample_header(m)
                            temp_context[table_title_k] = real_header if real_header else ''
                            if identity:
                                # 表格数据
                                data = self.get_unknown_sample_data(identity)
                                table_origin_header_data_mapping[jinja_header_value] = data
                        else:
                            temp_context[table_title_k] = self.TABLE_HEADER_MAP.get(table_title_k)
                            data = [f'样品{n+1}' for n in range(5)]
                            table_origin_header_data_mapping[jinja_header_value] = ['*', '*', '*', '*']
                
                    # 为每个表格生成行数据
                    if table_origin_header_data_mapping:
                        rows = self.columns_to_rows(table_origin_header_data_mapping)
                        rows_name = table_name + '_rows'
                        temp_context[rows_name] = rows
                return temp_context
        else:
            pass
        
        return temp_context  # 确保总是返回字典 

    async def get_quantitative_determination_photometric_measurement_from_database(self, doc_stream, task_id, origin_header_real_header_mapping={}, tpl_context={}):
        table_data_processor = TableDataProcessor(self.name)
        table_fields = table_data_processor.get_table_fields()
        if table_fields:
            if not hasattr(table_data_processor.handler_instance, 'HEADER_M_MAPPING'):
                raise APIException(error=ErrorCode.TABLE_DATA_HANDLER_NOT_FOUND)
            HEADER_M_MAPPING = table_data_processor.handler_instance.HEADER_M_MAPPING  # 每个私有文件都要声明的
            TABLE_NAME_MAP = table_data_processor.handler_instance.TABLE_NAME_MAP
            TEMPLATE_INFO_MAP = table_data_processor.handler_instance.TEMPLATE_INFO_MAP
            get_header_and_data_class_instance = GetHeaderAndData(task_id)
            await get_header_and_data_class_instance._get_all_headers()
            self.HEADER_M_MAPPING = HEADER_M_MAPPING
            temp_header_identity_map_list = []
            for origin_header, m in HEADER_M_MAPPING.items():
                real_header = None
                real_header, headerIdentity = self.get_real_header_from_endpoint_data(m)
                if not real_header:
                    _, real_header = await get_header_and_data_class_instance.get_header(m)
                    headerIdentity = get_header_and_data_class_instance.m_data_key
                if real_header:
                    # print('origin_header======================', origin_header)
                    bookmark_name = TABLE_NAME_MAP.get(origin_header) + '_' + origin_header
                    origin_header_real_header_mapping[bookmark_name] = real_header if real_header else origin_header # old_header 标签名字
                    """
                    identity: 用来查数据的标识符
                    real_header: 真正的表头。
                    m: 这个字段的标识
                    bookmark_name: 标签名字
                    tel_val: 表格循环的字段
                    比如 表格打上标签 "数据表格" 
                        bookmark_name=表头="数据表格_序号" 
                        tel_val="序号"
                    """
                    temp_header_identity_map_list.append({'identity': headerIdentity, 'real_header': real_header, 'm': m, 'bookmark_name': bookmark_name, 'tel_val': origin_header})
            # 把原来的表头替换成真实的表头

        
            get_header_and_data_class_instance.header_identity_map_list = temp_header_identity_map_list
            await get_header_and_data_class_instance.get_data()
            tel_val_data_list = get_header_and_data_class_instance.tel_val_data_list
            for template_name, template_info in TEMPLATE_INFO_MAP.items():
                if template_info.get('tpl_type') == 1:
                    # column_count = template_info.get('column_count')
                    # 每个单元格一个tpl对象，就是横向表格那种
                    TABLE_NAME_MAP_REVERSE = table_data_processor.handler_instance.TABLE_NAME_MAP_REVERSE
                    if TABLE_NAME_MAP_REVERSE.get(template_name):
                        for tel_val in TABLE_NAME_MAP_REVERSE.get(template_name):
                            if tel_val in tel_val_data_list:
                                tpl_val_data = tel_val_data_list.get(tel_val)
                                for i in range(len(tpl_val_data)):
                                    tpl_context[tel_val + '_' + str(i)] = tpl_val_data[i]
                if template_info.get('tpl_type') == 2:
                    TABLE_NAME_MAP_REVERSE = table_data_processor.handler_instance.TABLE_NAME_MAP_REVERSE
                    if TABLE_NAME_MAP_REVERSE.get(template_name):
                        new_tpl_val_data = {}
                        for tel_val in TABLE_NAME_MAP_REVERSE.get(template_name):
                            if tel_val in tel_val_data_list:
                                new_tpl_val_data[tel_val] = tel_val_data_list.get(tel_val)
                        rows = self.convert_tel_val_data_to_rows(new_tpl_val_data)
                        # 通过docxtpl循环语法渲染
                        tpl_context[template_name + '_rows'] = rows

            if hasattr(table_data_processor.handler_instance, 'get_grouped_context_for_template'):
                grouped_context = table_data_processor.handler_instance.get_grouped_context_for_template(tpl_context)
                return origin_header_real_header_mapping, grouped_context
            else:
                return origin_header_real_header_mapping, tpl_context

    
    def get_table_head_data_from_endpoint(self, task_id, analysis_id, all_display=0, digit=None):
        if all_display:
            all_display=True
        else:
            all_display=False
            
        if digit is not None:
            template_common_url = config.guide_template_table_head_data_url + '?taskId=' + urllib.parse.quote(str(task_id), safe='') + '&analysisId=' + urllib.parse.quote(str(analysis_id), safe='') + '&allDisplay=' + urllib.parse.quote(str(all_display), safe='') + '&digit=' + urllib.parse.quote(str(digit), safe='')
        else:
            template_common_url = config.guide_template_table_head_data_url + '?taskId=' + urllib.parse.quote(str(task_id), safe='') + '&analysisId=' + urllib.parse.quote(str(analysis_id), safe='') + '&allDisplay=' + urllib.parse.quote(str(all_display), safe='')
        try:
            response = requests.get(template_common_url)
            res_json = response.json()
            # print(res_json)
            return res_json.get('data')
        except requests.exceptions.Timeout:
            print(f"请求超时，跳过: {template_common_url}")
            return {}
        except Exception as e:
            print(f"请求异常，跳过: {e}")
            return {}

    def get_stand_real_header_from_endpoint_data(self, m):
        real_header = None
        stand_head_list = self.end_point_header_and_data.get('standHeadList')
        for stand_head_info in stand_head_list:
            if stand_head_info.get('identity') == m:
                real_header = stand_head_info.get("printTitle")
        return real_header or ''
    
    def get_stand_data_list(self, m):

        stand_head_list = self.end_point_header_and_data.get('standDataList')
        data = []
        for stand_data_info in stand_head_list:
            if stand_data_info.get(m):
                data.append(stand_data_info.get(m)) if stand_data_info.get(m) else data.append('')
        return data
        
                
    def get_unknown_sample_header(self, m):
        """
        未知样品:
            头: 
                1 去mapRelation list[dict] 里面取  paramIdentity=m 然后取headIdentity = identity 
                    真实表头 paramName
                    if paramName is None:
                        拿  identity  去 unknowHeadList 匹配 identity 取 printTitle 就是真实的表头
                2 去unknowHeadList 里面找
                    identity = m    printTitle就是真实的表头
            数据:
                unknowDataList 
                根据identity取数据
        """
        identity = None
        real_header = None
        if self.end_point_header_and_data:
            map_relation = self.end_point_header_and_data.get('mapRelation')
            for one_map_relation in map_relation:
                if one_map_relation.get('paramIdentity') == m:
                    identity = one_map_relation.get('headIdentity')
                    real_header = one_map_relation.get('paramName')
                    if not real_header:
                        unknown_header_list = self.end_point_header_and_data.get('unknowHeadList')
                        for one_unknown_header in unknown_header_list:
                            if one_unknown_header.get('identity') == identity:
                                real_header = one_unknown_header.get('printTitle')
        if identity is None and real_header is None:
            unknown_header_list = self.end_point_header_and_data.get('unknowHeadList')
            for one_unknown_header in unknown_header_list:
                if one_unknown_header.get('identity') == m:
                    real_header = one_unknown_header.get('printTitle')
                    identity = m
        return real_header, identity


    def get_unknown_sample_data(self, identity: str) -> List[str]:
        """
        未知样品:
            头: 
                1 去mapRelation list[dict] 里面取  paramIdentity=m 然后取headIdentity = identity 
                    真实表头 paramName
                    if paramName is None:
                        拿  identity  去 unknowHeadList 匹配 identity 取 printTitle 就是真实的表头
                2 去unknowHeadList 里面找
                    identity = m    printTitle就是真实的表头
            数据:
                unknowDataList 
                根据identity取数据
        """
        data = []
        for one_unknown_data in self.end_point_header_and_data.get('unknowDataList'):
            if one_unknown_data.get(identity) or one_unknown_data.get(identity) == '':
                data.append(one_unknown_data.get(identity)) if one_unknown_data.get(identity) else data.append('')
            else:
                data.append('')
        return data
    
    @staticmethod
    def columns_to_rows(data: Dict[str, List]) -> List[Dict]:
        """
        将 {'a':[1,2,3], 'b':[4,5,6], 'c':[7,8,9]}
        转成 [{'a':1,'b':4,'c':7}, {'a':2,'b':5,'c':8}, {'a':3,'b':6,'c':9}]
        """
        if not data:
            return []
        keys = list(data.keys())
        columns = [data[k] for k in keys]
        return [dict(zip(keys, row)) for row in zip(*columns)]
    
    # ==================== 表格样式辅助方法 ====================
    @staticmethod
    def _remove_outer_borders_only(table):
        """只移除表格最外圈边框，保留内部网格线"""
        tbl = table._tbl
        tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
        tblBorders = OxmlElement('w:tblBorders')
        # 外圈边框设为无
        for border_name in ('top', 'left', 'bottom', 'right'):
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'nil')
            tblBorders.append(border)
        # 内部网格线保留（single 黑色细线）
        for border_name in ('insideH', 'insideV'):
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')  # 0.5pt
            border.set(qn('w:color'), '000000')
            tblBorders.append(border)
        tblPr.append(tblBorders)
        if tbl.tblPr is None:
            tbl.insert(0, tblPr)

    @staticmethod
    def _set_cell_margins(table, margin=0):
        """设置表格所有单元格的内边距为 0（让内容贴边）"""
        tbl = table._tbl
        tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
        tblCellMar = OxmlElement('w:tblCellMar')
        for side in ('top', 'left', 'bottom', 'right'):
            node = OxmlElement(f'w:{side}')
            node.set(qn('w:w'), str(margin))
            node.set(qn('w:type'), 'dxa')
            tblCellMar.append(node)
        tblPr.append(tblCellMar)
        if tbl.tblPr is None:
            tbl.insert(0, tblPr)

    @staticmethod
    def _set_table_full_width(table):
        """设置表格宽度为 100%"""
        tbl = table._tbl
        tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
        tblW = OxmlElement('w:tblW')
        tblW.set(qn('w:w'), '5000')  # 5000 = 100% (pct)
        tblW.set(qn('w:type'), 'pct')
        tblPr.append(tblW)
        if tbl.tblPr is None:
            tbl.insert(0, tblPr)

    @staticmethod
    def _set_first_column_wider(table, first_col_pct=15):
        """设置第一列比其他列宽一些（作为表头列）"""
        if not table.rows:
            return
        
        num_cols = len(table.rows[0].cells)
        if num_cols <= 1:
            return
        
        # 计算其他列的百分比
        remaining_pct = 100 - first_col_pct
        other_col_pct = remaining_pct / (num_cols - 1)
        
        # 百分比转换为 pct 单位（50 = 1%，所以乘以 50）
        first_col_width = int(first_col_pct * 50)
        other_col_width = int(other_col_pct * 50)
        
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                # 移除已有的 tcW
                for existing in tcPr.findall(qn('w:tcW')):
                    tcPr.remove(existing)
                tcW = OxmlElement('w:tcW')
                if i == 0:
                    tcW.set(qn('w:w'), str(first_col_width))
                else:
                    tcW.set(qn('w:w'), str(other_col_width))
                tcW.set(qn('w:type'), 'pct')
                tcPr.append(tcW)
    
    # 光谱扫描 通用数据
    # async def get_spectral_scanning_common_context(self, tpl, image_base64, contexts):
    #     spectral_scanning_mapping = common_data_mapping.get(self.measure_type)
    #     for jinja_val, common_val_key in spectral_scanning_mapping.items():
    #         if jinja_val == '光谱扫描图片':
    #             contexts[jinja_val] = await self.decode_base64_image(tpl, image_base64)
    #         elif jinja_val == '扫描范围':
    #             scan_range = self.common_val.get('startWaveLength', '')  + '至' + self.common_val.get('endWaveLength', '')
    #             self.contexts[jinja_val] = scan_range
    #         else:
    #             self.contexts[jinja_val] = self.common_val.get(common_val_key) if self.common_val.get(common_val_key) else ''

    # # 时间扫描 通用数据
    # async def get_time_scanning_common_context(self, tpl, image_base64, contexts):
    #     time_scanning_mapping = common_data_mapping.get(self.measure_type)
    #     for jinja_val, common_val_key in time_scanning_mapping.items():
    #         if jinja_val != '时间扫描图片':
    #             self.contexts[jinja_val] = self.common_val.get(common_val_key) if self.common_val.get(common_val_key) else ''
    #         else:
    #             self.contexts[jinja_val] = await self.decode_base64_image(tpl, image_base64)
    
    # # 定量测定 通用数据
    # async def get_quantitative_determination_common_context(self, tpl, contexts):
    #     quantitative_determination = common_data_mapping.get(self.measure_type)
    #     if self.common_val:
    #         for jinja_val, common_val_key in quantitative_determination.items():
    #                 self.contexts[jinja_val] = self.common_val.get(common_val_key) if self.common_val.get(common_val_key) else ''
    #         # 处理图片
    #         curve_info = self.end_point_header_and_data.get('curveInfo')
    #         if curve_info:
    #             image_data_list = curve_info.get('dataList') or []
    #             image_stand_point_list = curve_info.get('standPoints') or []
    #             image_unknown_point_list = curve_info.get('unknowPoints') or []
    #             if image_data_list:
    #                 quantitative_image = await self.generate_quantitative_chart(
    #                     image_data_list, 
    #                     image_stand_point_list, 
    #                     image_unknown_point_list,
    #                     tpl
    #                 )
    #                 self.contexts['定量分析图片'] = quantitative_image if quantitative_image else ''
    #             else:
    #                 self.contexts['定量分析图片'] = ''
    #         else:
    #             self.contexts['定量分析图片'] = ''
    #         # k
    #         k_str = self.common_val.get('curveParams')
    #         if k_str:
    #             k_parts = (k_str.split(',') + [None] * 5)[:5]
                
    #             for i, value in enumerate(k_parts):
    #                 if value:
    #                     value = str(Decimal(value))
    #                     self.contexts[f'k{i}'] = f'k{i}: '
    #                     self.contexts[f'k{i}_v'] = value
    #         k_values = []
    #         if contexts.get('k0'):  # k0 非空
    #             k_values.append({'label': contexts['k0'], 'value': contexts['k0_v']})
    #         if contexts.get('k1'):
    #             k_values.append({'label': contexts['k1'], 'value': contexts['k1_v']})
    #         if contexts.get('k2'):
    #             k_values.append({'label': contexts['k2'], 'value': contexts['k2_v']})
    #         if contexts.get('k3'):
    #             k_values.append({'label': contexts['k3'], 'value': contexts['k3_v']})
    #         if contexts.get('k4'):
    #             k_values.append({'label': contexts['k4'], 'value': contexts['k4_v']})
    #         self.contexts['k_values'] = k_values
    
    # # 光度测量 通用数据
    # async def photometric_measurement_common_context(self, tpl, contexts):
    #     # 光度测量
    #     if self.common_val:
    #         photometric_measurement = common_data_mapping.get(self.measure_type)
    #         for jinja_val, common_val_key in photometric_measurement.items():
    #             self.contexts[jinja_val] = self.common_val.get(common_val_key) if self.common_val.get(common_val_key) else ''
    
    # 光谱扫描 表格数据
    async def get_spectral_scanning_table_context(self):
        table_data = self.handle_spectral_scanning_table(self.end_point_header_and_data)
        # 兼容旧模板（仍可用 Jinja 循环渲染）
        self.contexts['table_data'] = table_data
        # 新方式：预先用 python-docx 生成表格，再通过 docxtpl subdoc 插入（只改光谱）
    
    # 时间扫描 表格数据
    async def get_time_scanning_table_context(self):
        table_data = self.handle_time_scanning_table(self.end_point_header_and_data)
        self.contexts['table_data'] = table_data
    
    # 定量测定 表格数据
    async def get_quantitative_determination_table_context(self):
        if self.fake:
            knows_header, knows_data_list = self.get_quantitative_determination_standard_table_fake()
            knows_table_data = {'headers': knows_header, 'row_data': knows_data_list}
            self.contexts['knows_table_data'] = knows_table_data
            unknown_header, unknown_data_list = self.get_quantitative_determination_unknown_table_fake()
            unknown_table_data = {'headers': unknown_header, 'row_data': unknown_data_list}
            self.contexts['unknown_table_data'] = unknown_table_data
        else:
            if self.end_point_header_and_data:
                # 标准样品表格
                standard_header_list = self.end_point_header_and_data.get('standHeadList')
                standard_data_list = self.end_point_header_and_data.get('standDataList')
                knows_header, knows_data_list = self.get_quantitative_determination_standard_table(standard_header_list, standard_data_list)
                knows_table_data = {'headers': knows_header, 'row_data': knows_data_list}
                self.contexts['knows_table_data'] = knows_table_data

                # abs_identity = None
                # same_number_identity = 'sampleNumber'
                # for header in standard_header_list:
                #     if header.get('locate') == 'measure' and 'Abs' in header.get('name') and '-' in header.get('identity'):
                #         abs_identity = header.get('identity')

                # contexts['波长'] = str(abs_identity.split('-')[0]) if abs_identity and '-' in abs_identity else ''

                # 未知样品表格
                unknown_data_list = self.end_point_header_and_data.get('unknowDataList')
                unknown_header_list = self.end_point_header_and_data.get('unknowHeadList')
                unknown_header, unknown_data_list = self.get_quantitative_determination_unknown_table(unknown_header_list, unknown_data_list)
                unknown_table_data = {'headers': unknown_header, 'row_data': unknown_data_list}
                self.contexts['unknown_table_data'] = unknown_table_data
            else:
                self.contexts['knows_table_data'] = {'headers': [], 'row_data': []}
                self.contexts['unknown_table_data'] = {'headers': [], 'row_data': []}
    
    # 光度测量 表格数据
    async def get_photometric_measurement_table_context(self):
        if self.fake:
            header = ['编号', '带宽', 'h2', 'h3', 'h4']
            row_data = []
            for i in range(len(header)):
                row = [i+1, 200 + i*10]
                for j in range(len(header)-len(row)):
                    row.append('--' if (i+j) % 2 == 0 else '**')
                row_data.append(row)
            self.contexts['table_data'] = {'headers': header, 'row_data': row_data}
        else:
            if self.end_point_header_and_data:
                header_list = self.end_point_header_and_data.get('headList')
                data_list = self.end_point_header_and_data.get('dataList')
                headers, row_data = self.photometric_measurement_table_data(header_list, data_list)
                table_data = {'headers': headers, 'row_data': row_data}
                self.contexts['table_data'] = table_data
            else:
                self.contexts['table_data'] = {'headers': [], 'row_data': []}
    
    # 光谱扫描
    def handle_spectral_scanning_table(self, end_point_header_and_data):
        if self.fake:
            header = ['编号', '波长', '样品1', '峰谷', '样品2', '峰谷', '样品3', '峰谷']
            row_data = []
            for i in range(10):
                row = [i+1, 200 + i*10]
                for j in range(3):
                    row.append(round(0.1 * (i+1) * (j+1), 3))
                    row.append('峰' if (i+j) % 2 == 0 else '谷')
                row_data.append(row)
            return {'header': header, 'row_data': row_data}
        else:
            head_list = end_point_header_and_data.get('headList')
            data_list = end_point_header_and_data.get('dataList')
            if self.single_sample:
                identity_list = []
                header = []
                row_data = []
                if head_list:
                    for head_mapping in head_list:
                        header.append(head_mapping.get('name'))
                        identity_list.append(head_mapping.get('identity'))
                    for data in data_list:
                        row = []
                        for index, identity in enumerate(identity_list):
                            if data.get(identity):
                                row.append(data.get(identity))
                            else:
                                row.append('')
                        row_data.append(row)
            else:
                sample_num = 0 # 样品数量
                sample_name = []
                row_data = []
                if head_list:
                    for i in head_list:
                        if i.get('identity') == 'sampleNumber':
                            sample_num += 1
                            sample_name.append(i.get('name'))
                    header = [h.get("name") for h in head_list]
                    # header = ['编号', '波长']
                    # for name in sample_name:
                    #     header.append(name)
                    #     header.append('峰谷')
                    for wave_length_data_mapping in data_list:
                        row = []
                        for wave_length, data in wave_length_data_mapping.items():
                            row.append(data.get('id'))
                            row.append(wave_length)
                            for index, sample_num in enumerate(data.get('sampleNumber')):
                                row.append(sample_num)
                                row.append(data.get('peakValley')[index])
                        row_data.append(row)
                else:
                    row_data = []

        return {'header': header, 'row_data': row_data}
    
    # 时间扫描
    def handle_time_scanning_table(self, end_point_header_and_data):
        if self.fake:
            header = ['编号', '带宽', 'h2', 'h3', 'h4']
            row_data = []
            for i in range(len(header)):
                row = [i+1, 200 + i*10]
                for j in range(len(header)-len(row)):
                    row.append('--' if (i+j) % 2 == 0 else '**')
                row_data.append(row)
            return {'header': header, 'row_data': row_data}
        else:
            head_list = end_point_header_and_data.get('headList')
            data_list = end_point_header_and_data.get('dataList')
            if self.single_sample:
                identity_list = []
                header = []
                for head_mapping in head_list:
                    header.append(head_mapping.get('name'))
                    identity_list.append(head_mapping.get('identity'))
                row_data = []
                for data in data_list:
                    row = []
                    for index, identity in enumerate(identity_list):
                        if data.get(identity):
                            row.append(data.get(identity))
                        else:
                            row.append('')
                    row_data.append(row)
            else:
                sample_num = 0 # 样品数量
                sample_name = []
                for i in head_list:
                    if i.get('identity') == 'sampleNumber':
                        sample_num += 1
                        sample_name.append(i.get('name'))
                header = [h.get("name") for h in head_list]
                row_data = []
                for wave_length_data_mapping in data_list:
                    row = []
                    for time, data in wave_length_data_mapping.items():
                        row.append(data.get('id'))
                        row.append(time)
                        for index, sample_num in enumerate(data.get('sampleNumber')):
                            row.append(sample_num)
                            if data.get('peakValley')[index]:
                                row.append(data.get('peakValley')[index])
                            else:
                                row.append('')
                    row_data.append(row)

        return {'header': header, 'row_data': row_data}
        

    # 定量已知样
    def get_quantitative_determination_standard_table(self, standard_header_list, standard_data_list):
        if self.fake:
            pass
        else:
            name_identity_mapping = {header.get('name'): header.get('identity') for header in standard_header_list}
            headers = []
            for h in standard_header_list:
                if h.get('printTitle'):
                    headers.append(h.get('printTitle'))
                else:
                    headers.append(h.get('name'))
            row_data = []
            for data in standard_data_list:
                row = []
                for _, identity in name_identity_mapping.items():
                    row_value = data.get(identity)
                    if row_value:
                        row.append(row_value)
                    else:
                        row.append('')
                row_data.append(row)
                if data.get('children', []):
                    for children_data in data.get('children', []):
                        row = []
                        for _, identity in name_identity_mapping.items():
                            row_value = children_data.get(identity)
                            if row_value:
                                row.append(row_value)
                            else:
                                row.append('')
                        row_data.append(row)
        return headers, row_data

    def get_quantitative_determination_standard_table_fake(self):
        header = ['样品编号', '浓度', '已知样1', '已知样2','吸光度', '备注']
        row_data = []
        for i in range(5):
            row = [f'STD-{i+1}', round(10.0 * (i+1), 2), round(0.5 * (i+1), 3), round(0.6 * (i+1), 3), round(0.05 * (i+1), 3), '无']
            row_data.append(row)
        return header, row_data
    
    # 定量未知样
    def get_quantitative_determination_unknown_table(self, unknown_header_list, unknown_data_list):
        name_identity_mapping = {header.get('name'): header.get('identity') for header in unknown_header_list}
        headers = []
        for h in unknown_header_list:
            if h.get('printTitle'):
                headers.append(h.get('printTitle'))
            else:
                headers.append(h.get('name'))
        row_data = []
        for data in unknown_data_list:
            row = []
            for _, identity in name_identity_mapping.items():
                row_value = data.get(identity)
                if row_value:
                    row.append(row_value)
                else:
                    row.append('')
            row_data.append(row)
            if data.get('children', []):
                for children_data in data.get('children', []):
                    row = []
                    for _, identity in name_identity_mapping.items():
                        row_value = children_data.get(identity)
                        if row_value:
                            row.append(row_value)
                        else:
                            row.append('')
                    row_data.append(row)
        return headers, row_data

    def get_quantitative_determination_unknown_table_fake(self):
        header = ['样品编号', '未知样1', '未知样2','吸光度', '备注']
        row_data = []
        for i in range(5):
            row = [f'UNK-{i+1}', round(0.7 * (i+1), 3), round(0.8 * (i+1), 3), round(0.06 * (i+1), 3), '无']
            row_data.append(row)
        return header, row_data
    
    # 光度测量
    def photometric_measurement_table_data(self, header_list, data_list):
        if self.fake:
            pass
        else:
            name_identity_mapping = {header.get('name'): header.get('identity') for header in header_list}
            headers = []
            for h in header_list:
                if h.get('printTitle'):
                    headers.append(h.get('printTitle'))
                else:
                    headers.append(h.get('name'))
            row_data = []
            for data in data_list:
                row = []
                for _, identity in name_identity_mapping.items():
                    row_value = data.get(identity)
                    if row_value:
                        row.append(row_value)
                    else:
                        row.append('')
                row_data.append(row)
                if data.get('children', []):
                    for children_data in data.get('children', []):
                        row = []
                        for name in headers:
                            identity = name_identity_mapping.get(name)
                            if identity:
                                row_value = children_data.get(identity)
                                if row_value:
                                    row.append(row_value)
                                else:
                                    row.append('')
                        row_data.append(row)
        return headers, row_data


