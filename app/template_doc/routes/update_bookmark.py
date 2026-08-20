# 标准库导入
import base64
import io
from io import BytesIO
import os
import subprocess
import tempfile
import time
from typing import List, Optional, Union
from urllib.parse import quote

# 第三方库导入
from docx import Document
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from tortoise import connections

# 本地应用导入
from ..model.template_bookmark import TemplateBookmark
from ..model.template_doc import TemplateDoc
from .create_template_view import create_template
from app.classic_template_doc.routes.classic_create_template_view import create_template as classic_create_template
from config import config
from constants.template_constant import TemplatePageNumType, MeasureMode, TemplateFootAndHeaderType
from exceptions.command_exception import APIException
from exceptions.error import ErrorCode
from utils.EditDocxXML import fill_word_bookmark
from utils.classic_html_builder import ClassicHtmlBuilder
from utils.doc_converter import get_libreoffice_command, get_libreoffice_subprocess_kwargs, explain_libreoffice_returncode, get_libreoffice_timeout_seconds
from utils.fille_template_utils import FillDataUtils
from utils.handle_header_footer import set_header_footer
from utils.log_decorator import log_response
from utils.template_utils import get_doc_obj_by_template_id, get_title_mapping

router = APIRouter()

class BookmarkItem(BaseModel):
    bm_id: Optional[int] = Field(None, description="书签id，更新时传入")
    desc_t: str = Field(None, description="time_mapping 的key")
    font: Optional[str] = Field(None, description="字体")
    font_size: Optional[Union[int, float]] = Field(None, description="字号")
    value: Optional[str] = Field(None, description="书签内容")


class FootHeaderItem(BaseModel):
    value: Optional[str] = Field(None, description="书签内容")
    font: Optional[str] = Field(None, description="字体")
    font_size: Optional[Union[int, float]] = Field(None, description="字号")
    position: Optional[str] = Field(None, description="0,1,2分别对应右,中,左, 和 TemplateFootAndHeaderType保持一致")
    type: Optional[int] = Field(None, description="0=关联title_mapping字段 1=自定义文本")



class InputSchema(BaseModel):
    template_id: int = Field(..., description='模板id')
    bookmark_list: List[BookmarkItem] = Field(None, description='标签名称列表')
    title_list: List[BookmarkItem] = Field(None, description='报告名称')
    operation_type: int = Field(None, description='操作类型 0: 预览 1:另存为 2: 保存')
    name: str = Field(None, description='模板名称 当另存为的时候传入')
    remark: Optional[BookmarkItem]= Field(None, description='备注')
    project_name: Optional[BookmarkItem] = Field(None, description='动态项目名称')
    page_num_type: str = Field(TemplatePageNumType.no_page.value, description='页码类型')
    measure_mode: int = Field(None, description='测量模式 1:默认 2:仅测量标签位置不修改标签字体大小')
    margin: Optional[str] = Field(None, description='动态边距(mm)')
    orientation: Optional[int] = Field(None, description='横向/纵向 0=纵向 1=横向')
    line_height: Optional[str] = Field(None, description='行边距')
    left_header: Optional[FootHeaderItem] = Field(None, description='左页眉')
    mid_header: Optional[FootHeaderItem] = Field(None, description='中页眉')
    right_header: Optional[FootHeaderItem] = Field(None, description='右页眉')
    left_footer: Optional[FootHeaderItem] = Field(None, description='左页脚')
    mid_footer: Optional[FootHeaderItem] = Field(None, description='中页脚')
    right_footer: Optional[FootHeaderItem] = Field(None, description='右页脚')


class DataSchema(BaseModel):
    file_base64: str = Field(..., description='文件base64')


def _desc_to_position(desc: Optional[str]) -> str:
    """将 TemplateFootAndHeaderType 描述文本映射为 CSS position 值（'left'/'center'/'right'）。"""
    if not desc:
        return 'center'
    _VALUE_TO_POS = {
        TemplateFootAndHeaderType.left.value: 'left',
        TemplateFootAndHeaderType.center.value: 'center',
        TemplateFootAndHeaderType.right.value: 'right',
    }
    decs_value_map = TemplateFootAndHeaderType.generate_decs_value_map()
    value = decs_value_map.get(desc)
    return _VALUE_TO_POS.get(value, 'center')


def _collect_header_footer_configs(data: InputSchema, side: str) -> dict:
    """把前端传的 6 个 FootHeaderItem 转成配置 dict，供 resolve_header_footer_items 使用。"""
    from utils.template_utils import normalize_position
    side_items = {
        'header': ['left_header', 'mid_header', 'right_header'],
        'footer': ['left_footer', 'mid_footer', 'right_footer'],
    }
    configs = {}
    for key in side_items.get(side, []):
        item = getattr(data, key, None)
        if not item:
            continue
        configs[key] = {
            'type': item.type if item.type is not None else 1,
            'value': item.value,
            'font': item.font,
            'font_size': item.font_size,
            'position': normalize_position(item.position),
        }
    return configs


def _set_header_footer_three_parts(doc, hdr_items=None, ftr_items=None):
    """
    向导模式（DOCX）：左中右三段页眉页脚。
    用一行三列的无边框表格实现左中右分布（比制表符方案可靠），缺位留空。
    """
    from docx.shared import Pt
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn as _qn
    from docx.oxml import OxmlElement

    def _fill(container, items, default_size):
        # 先删除所有旧段落，再加占位段落（保证 add_table 有锚点），
        # 表格插入后再删除占位段落，页眉页脚只剩无边框表格
        for para in container.paragraphs[:]:
            para._element.getparent().remove(para._element)
        container.add_paragraph()

        # 计算可用宽度：页面宽度 - 左右边距（边距可能为 None，给默认值）
        try:
            from docx.shared import Cm
            _section = container.part.document.sections[0]
            _page_w = _section.page_width if _section.page_width is not None else Cm(21)
            _left_m = _section.left_margin if _section.left_margin is not None else Cm(2.54)
            _right_m = _section.right_margin if _section.right_margin is not None else Cm(2.54)
            _usable = _page_w - _left_m - _right_m
        except Exception:
            from docx.shared import Cm
            _usable = Cm(15.9)
        if _usable is None or _usable <= 0:
            from docx.shared import Cm
            _usable = Cm(15.9)

        _col_w = _usable // 3
        widths = [_col_w, _col_w, _usable - 2 * _col_w]

        table = container.add_table(rows=1, cols=3, width=_usable)
        table.autofit = False
        for i in range(3):
            table.columns[i].width = widths[i]
            table.rows[0].cells[i].width = widths[i]

        slots = {it.get('position', 'left'): it for it in items}
        _align_map = {
            'left': WD_PARAGRAPH_ALIGNMENT.LEFT,
            'center': WD_PARAGRAPH_ALIGNMENT.CENTER,
            'right': WD_PARAGRAPH_ALIGNMENT.RIGHT,
        }
        for i, pos in enumerate(('left', 'center', 'right')):
            cell = table.rows[0].cells[i]
            it = slots.get(pos)
            text = it.get('value', '') if it else ''
            p = cell.paragraphs[0]
            p.alignment = _align_map[pos]
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.left_indent = None
            p.paragraph_format.right_indent = None
            if not text:
                continue
            run = p.add_run(text)
            if it:
                _font = it.get('font') or '宋体'
                run.font.name = _font
                if it.get('font_size'):
                    run.font.size = Pt(it['font_size'])
                _rPr = run._element.get_or_add_rPr()
                _rFonts = _rPr.get_or_add_rFonts()
                _rFonts.set(_qn('w:eastAsia'), _font)

        # 去表格边框和单元格边距
        _tbl = table._tbl
        _tblPr = _tbl.tblPr
        _borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            _el = OxmlElement(f'w:{edge}')
            _el.set(_qn('w:val'), 'none')
            _el.set(_qn('w:sz'), '0')
            _el.set(_qn('w:space'), '0')
            _borders.append(_el)
        _tblPr.append(_borders)
        _cell_mar = OxmlElement('w:tblCellMar')
        for side in ('top', 'left', 'bottom', 'right'):
            _el = OxmlElement(f'w:{side}')
            _el.set(_qn('w:w'), '0')
            _el.set(_qn('w:type'), 'dxa')
            _cell_mar.append(_el)
        _tblPr.append(_cell_mar)

        # 删除表格后面的占位段落
        for para in container.paragraphs[:]:
            para._element.getparent().remove(para._element)

    for section in doc.sections:
        if hdr_items:
            header = section.header
            header.is_linked_to_previous = False
            _fill(header, hdr_items, 10)
        if ftr_items:
            footer = section.footer
            footer.is_linked_to_previous = False
            _fill(footer, ftr_items, 9)


@router.post('/update_bookmark')
@log_response
async def update_bookmark(data: InputSchema):
    if data.measure_mode == MeasureMode.classic.value:
        class_db_con = connections.get('classic_db')
        template = await TemplateDoc.get_or_none(id=data.template_id).using_db(class_db_con)
    else:
        template = await TemplateDoc.get_or_none(id=data.template_id)
    if not template:
        raise APIException(ErrorCode.TEMPLATE_NOT_FOUND)
    template_name = template.name
    origin_title_mapping = get_title_mapping(template_name, measure_mode=data.measure_mode)
    operation_type = data.operation_type
    input_title_mapping = data.bookmark_list
    title_list = data.title_list
    new_title_mapping = {}
    for item in input_title_mapping:
        new_title_mapping[item.desc_t] = {
            'name': item.value,
            'font': item.font,
            'font_size': item.font_size
        }
    if operation_type == 1:  # 另存为
        new_name = data.name
        if not new_name:
            raise APIException(ErrorCode.TEMPLATE_NAME_EMPTY)
        # 合并标题映射
        merged_title_mapping = {k: new_title_mapping[k] for k in origin_title_mapping if k in new_title_mapping}
        content_base64 = template.template_content
        content = base64.b64decode(content_base64)
        thumbnail_base64 = template.thumbnail_base64

        is_classic_save = data.measure_mode == MeasureMode.classic.value
        # 收集左中右页眉页脚配置（写入 .py 配置的 extra_info）
        _hdr_configs = _collect_header_footer_configs(data, 'header')
        _ftr_configs = _collect_header_footer_configs(data, 'footer')

        if data.measure_mode == MeasureMode.classic.value:
            await classic_create_template(new_name, content_base64, template_type=template.type, thumbnail_base64=thumbnail_base64, experiment_type=template.experiment_type)
            # 创建 .py 配置文件（从原模板复制，更新 title_mapping、标题、页眉页脚、页码）
            page_num = None
            if data.page_num_type:
                _VALUE_TO_POS = {
                    TemplatePageNumType.left_page.value: 'left',
                    TemplatePageNumType.center_page.value: 'center',
                    TemplatePageNumType.right_page.value: 'right',
                }
                decs_value_map = TemplatePageNumType.generate_decs_value_map()
                value = decs_value_map.get(data.page_num_type)
                if value is not None and value != TemplatePageNumType.no_page.value:
                    pos = _VALUE_TO_POS.get(value)
                    page_num = {'position': pos, 'font': '宋体', 'font_size': 9}
                else:
                    page_num = {'position': '', 'font': '宋体', 'font_size': 9}  # 不展示页码时 position 为空
            _save_classic_config_file(template_name, new_name, merged_title_mapping, _hdr_configs, _ftr_configs, title_list, page_num, data.margin, data.orientation, data.line_height)
        else:
            # 向导 HTML 模板另存为：复制 HTML 模板 + 复制更新 .py 配置
            from app.template_doc.routes.create_template_view import create_guide_html_template
            await create_guide_html_template(
                data.name, content_base64, content=content,
                template_type=template.type,
                industry_type=template.industry_type,
                experiment_type=template.experiment_type,
            )
            _save_classic_config_file(
                template_name, new_name, merged_title_mapping, _hdr_configs, _ftr_configs,
                title_list, None, None, None, None,
                py_dir=str(config.structure_path),
            )
        return {'code': 200, 'message': '模板另存为成功', 'data': {'file_base64': ''}}

    else:  # 预览

        # ---- 经典模式：从 DB 取 HTML 模板，Playwright 渲染 PDF ----
        if data.measure_mode == MeasureMode.classic.value:

            # 从 title_mapping 构建 contexts（_t = 标签名，值留空，_fs/_font = 字号/字体）
            contexts = {}
            for key, info in new_title_mapping.items():
                if key.endswith('_t'):
                    base = key[:-2]
                    contexts[key] = info.get('name', '')           # 标签
                    contexts[base] = ''                             # 值留空
                    if info.get('font_size'):
                        contexts[base + '_fs'] = info['font_size']
                    if info.get('font'):
                        contexts[base + '_font'] = info['font']

            if data.remark and data.remark.value:
                contexts['备注'] = data.remark.value

            # 标题列表（前端传过来的修改后名称）
            if title_list:
                contexts['__title_list__'] = [
                    {'name': getattr(t, 'desc_t', ''), 'value': t.value, 'font': t.font, 'font_size': t.font_size}
                    for t in title_list
                ]

            # 从配置文件读取标题、页码、页眉页脚（前端传的会覆盖）
            try:
                from utils.template_utils import import_template_module, build_header_footer_context
                mod = import_template_module(config.classic_structure_module, template_name)
                if hasattr(mod, 'extra_info'):
                    ei = mod.extra_info
                    # 页码：前端 page_num_type 优先，否则用配置文件
                    if 'page_num' in ei:
                        contexts['__page_num__'] = dict(ei['page_num'])
                    # 配置文件默认值（会被前端覆盖）
                    contexts.setdefault('__title_list__', ei.get('title_list', []))
                    contexts.setdefault('__header__', build_header_footer_context(ei, 'header', contexts=contexts))
                    contexts.setdefault('__footer__', build_header_footer_context(ei, 'footer', contexts=contexts))
            except Exception:
                pass

            # 前端传的覆盖配置文件（使用 TemplatePageNumType 常量映射）
            if data.page_num_type:
                _VALUE_TO_POS = {
                    TemplatePageNumType.left_page.value: 'left',
                    TemplatePageNumType.center_page.value: 'center',
                    TemplatePageNumType.right_page.value: 'right',
                }
                decs_value_map = TemplatePageNumType.generate_decs_value_map()
                value = decs_value_map.get(data.page_num_type)
                if value is not None and value != TemplatePageNumType.no_page.value:
                    pos = _VALUE_TO_POS.get(value, 'center')
                    contexts.setdefault('__page_num__', {})['position'] = pos
                else:
                    # 不展示页码：清除配置文件中的默认值
                    contexts.pop('__page_num__', None)

            # 左中右页眉页脚（前端传的覆盖配置文件）
            from utils.template_utils import resolve_header_footer_items
            _hdr_configs = _collect_header_footer_configs(data, 'header')
            _hdr_items = resolve_header_footer_items(_hdr_configs, 'header', contexts=contexts)
            if _hdr_items:
                contexts['__header__'] = {'items': _hdr_items}
            _ftr_configs = _collect_header_footer_configs(data, 'footer')
            _ftr_items = resolve_header_footer_items(_ftr_configs, 'footer', contexts=contexts)
            if _ftr_items:
                contexts['__footer__'] = {'items': _ftr_items}
            # 边距 / 横向纵向 / 行边距（前端传的覆盖配置文件）
            if data.margin is not None:
                contexts['__margin__'] = {'value': data.margin}
            if data.orientation is not None:
                contexts['__orientation__'] = {'value': data.orientation}
            if data.line_height is not None:
                contexts['__line_height__'] = {'value': data.line_height}

            # 计算各区块是否有字段（用于隐藏空区块标题）
            for _flag, _keys in [
                ('__has_instrument__',   ['仪器型号_t', '仪器序号_t', '光谱带宽_t', '激发狭缝_t', '发射狭缝_t', '激发滤光片_t', '发射滤光片_t', '灵敏度_t', '发射波长_t']),
                ('__has_file_info__',    ['数据文件_t', '分析时间_t', '数据类型_t', '方法名称_t']),
                ('__has_calibration__',  ['曲线方程_t', '方程次数_t', '公式_t', '校正方法_t', '自动切换样品池_t', '零点插入_t', '测量波长_t', '强制过零点_t', '扫描模式_t', '激发波长_t', '发射波长_t', '重复测量_t']),
                ('__has_spectrum_scan__',['扫描范围_t', '光度模式_t', '增益_t', '波长间隔_t', '扫描间隔_t', '扫描速度_t', '能量灯_t', '数据文件_t', '分析时间_t', '数据类型_t', '方法名称_t']),
                ('__has_time_scan__',    ['光度模式_t', '增益_t', '时间单位_t', '扫描时间_t', '时间间隔_t', '扫描速度_t', '能量灯_t', '采样数_t', '测量波长_t', '数据文件_t', '分析时间_t', '数据类型_t', '方法名称_t']),
                ('__has_analysis_record__',['分析员_t', '审核员_t', '备注_t', '注释_t']),
            ]:
                contexts[_flag] = any(contexts.get(k) for k in _keys)

            # 从 DB 获取 HTML 模板字符串
            html_template_str = _get_html_from_template(template)
            builder = ClassicHtmlBuilder(template_html=html_template_str)
            try:
                pdf_bytes = builder.build_pdf(contexts)
            except Exception as e:
                print(f"ERROR: ClassicHtmlBuilder 渲染失败: {e}")
                import traceback
                traceback.print_exc()
                raise APIException(ErrorCode.CONVERT_PDF_FAILED)

            filename = f"{template_name}_filled.pdf"
            safe_name = quote(filename)
            return StreamingResponse(
                BytesIO(pdf_bytes),
                media_type="application/pdf",
                headers={
                    "Content-Disposition": f"attachment; filename*=UTF-8''{safe_name}",
                    "Content-Length": str(len(pdf_bytes)),
                },
            )

        # ---- 向导模式（HTML）：Jinja2 + Playwright 渲染 ----

        # 从 DB 获取 HTML 模板字符串
        html_template_str = _get_html_from_template(template)
        if not html_template_str.strip():
            print(f"[向导模式] 模板 {template.name}(id={template.id}) 的 template_content 为空或非 HTML")
            raise APIException(ErrorCode.TEMPLATE_NOT_FOUND, "向导模板内容为空，请重新上传 HTML 模板")
        print(f"[向导模式] 模板 id={template.id} name={template.name} → HTML 渲染")

        # contexts：title_mapping 标签 + 值 + 表格数据
        fill_data_utils = FillDataUtils(template_name, fake=True, new_title_mapping=new_title_mapping, template=template, structure_path=config.structure_module)
        contexts = await fill_data_utils.build_html_contexts(
            remark=data.remark.value if data.remark else None,
        )

        # 前端传的标签名覆盖
        for key, info in new_title_mapping.items():
            if key.endswith('_t') and info.get('name'):
                contexts[key] = info['name']
                base = key[:-2]
                if info.get('font_size'):
                    contexts[base + '_fs'] = info['font_size']
                if info.get('font'):
                    contexts[base + '_font'] = info['font']

        # 标题列表
        if title_list:
            contexts['__title_list__'] = [
                {'name': getattr(t, 'desc_t', ''), 'value': t.value, 'font': t.font, 'font_size': t.font_size}
                for t in title_list
            ]

        # 页码类型（desc 文本 → position）
        if data.page_num_type:
            _VALUE_TO_POS = {
                TemplatePageNumType.left_page.value: 'left',
                TemplatePageNumType.center_page.value: 'center',
                TemplatePageNumType.right_page.value: 'right',
            }
            decs_value_map = TemplatePageNumType.generate_decs_value_map()
            value = decs_value_map.get(data.page_num_type)
            if value is not None and value != TemplatePageNumType.no_page.value:
                contexts['__page_num__'] = {'position': _VALUE_TO_POS.get(value, 'center'), 'font': '宋体', 'font_size': 9}

        # 左中右页眉页脚
        from utils.template_utils import resolve_header_footer_items
        _hdr_items = resolve_header_footer_items(_collect_header_footer_configs(data, 'header'), 'header', contexts=contexts)
        if _hdr_items:
            contexts['__header__'] = {'items': _hdr_items}
        _ftr_items = resolve_header_footer_items(_collect_header_footer_configs(data, 'footer'), 'footer', contexts=contexts)
        if _ftr_items:
            contexts['__footer__'] = {'items': _ftr_items}

        # 边距 / 横向纵向 / 行边距
        if data.margin is not None:
            contexts['__margin__'] = {'value': data.margin}
        if data.orientation is not None:
            contexts['__orientation__'] = {'value': data.orientation}
        if data.line_height is not None:
            contexts['__line_height__'] = {'value': data.line_height}

        builder = ClassicHtmlBuilder(template_html=html_template_str)
        try:
            pdf_bytes = builder.build_pdf(contexts)
        except Exception as e:
            print(f"ERROR: ClassicHtmlBuilder 渲染失败: {e}")
            import traceback
            traceback.print_exc()
            raise APIException(ErrorCode.CONVERT_PDF_FAILED)

        filename = f"{template_name}_filled.pdf"
        safe_name = quote(filename)
        return StreamingResponse(
            BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{safe_name}",
                "Content-Length": str(len(pdf_bytes)),
            },
        )

async def insert_title(title_list, doc, template_id, measure_mode):
    title_list = title_list
    title_list_id = [item.bm_id for item in title_list]
    title_list_id_info_mapping = {item.bm_id: item for item in title_list}
    title_bookmark_instance = await TemplateBookmark.filter(template_doc_id=template_id, id__in=title_list_id, is_del = 0)
    for bookmark in title_bookmark_instance:
        name = bookmark.name
        value = title_list_id_info_mapping[bookmark.id].value
        bookmark_type = bookmark.type
        doc_stream = fill_word_bookmark(doc, name, value, bookmark_type, font_size=title_list_id_info_mapping[bookmark.id].font_size)
    return doc_stream

def apply_font_and_size(doc, font=None, font_size=None):
    """
    统一修改文档字体和字号。
    - font:      对所有 run 生效（中英文字体均修改）
    - font_size: 跳过标题段落（样式名含 'heading' 或 '标题'）
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    HEADING_KEYWORDS = ('heading', '标题')

    def is_heading(paragraph):
        try:
            return any(kw in paragraph.style.name.lower() for kw in HEADING_KEYWORDS)
        except Exception:
            return False

    def process_run(r_elem, skip_size):
        rPr = r_elem.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            r_elem.insert(0, rPr)

        if font:
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.insert(0, rFonts)
            rFonts.set(qn('w:ascii'), font)
            rFonts.set(qn('w:hAnsi'), font)
            rFonts.set(qn('w:eastAsia'), font)
            rFonts.set(qn('w:cs'), font)

        if font_size and not skip_size:
            sz_val = str(int(float(font_size) * 2))  # Word XML 单位为半磅
            for tag_name in ('w:sz', 'w:szCs'):
                el = rPr.find(qn(tag_name))
                if el is None:
                    el = OxmlElement(tag_name)
                    rPr.append(el)
                el.set(qn('w:val'), sz_val)

    # 收集正文段落 + 表格内段落
    all_paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)

    for para in all_paragraphs:
        skip_size = is_heading(para)
        for run in para.runs:
            process_run(run._r, skip_size)


def file_stream_to_base64(file_stream):
    doc_raw_bytes = file_stream.read()
    base64_str = base64.b64encode(doc_raw_bytes).decode("utf-8")
    return base64_str, doc_raw_bytes


def convert_docx_to_images_base64(file_stream) -> List[str]:
    """
    将Word文档转换为图片列表，返回base64编码列表
    适用于预览功能，保持格式完整性
    """
    doc_raw_bytes = file_stream.read()
    
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "input.docx")
        
        # 保存docx文件
        with open(docx_path, "wb") as f:
            f.write(doc_raw_bytes)
        
        try:
            # 使用LibreOffice转换为图片（PNG格式保持高质量）
            cmd=get_libreoffice_command(tmpdir,docx_path,'png')
            result = subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=get_libreoffice_timeout_seconds(), **get_libreoffice_subprocess_kwargs(cmd))
            
            # 收集所有生成的图片并转换为base64
            image_base64_list = []
            png_files = sorted([f for f in os.listdir(tmpdir) if f.endswith('.png')])
            
            for filename in png_files:
                img_path = os.path.join(tmpdir, filename)
                with open(img_path, "rb") as img_file:
                    img_base64 = base64.b64encode(img_file.read()).decode('utf-8')
                    image_base64_list.append(f"data:image/png;base64,{img_base64}")
            
            return image_base64_list
            
        except subprocess.CalledProcessError as e:
            detail = explain_libreoffice_returncode(e.returncode)
            if detail:
                print(f"LibreOffice诊断: {detail}")
            print(f"LibreOffice转换失败: {e}")
            return []
        except Exception as e:
            print(f"转换过程出错: {e}")
            return []

def convert_docx_to_pdf(docx_bytes: bytes) -> bytes:
    """
    使用 LibreOffice 将 docx 文件转换为 pdf 并返回 pdf 的字节内容。
    
    参数：
        docx_bytes: Word 文档的二进制内容（例如从 BytesIO 或 UploadFile.read() 得到）

    返回：
        转换后的 PDF 文件内容（bytes）
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "input.docx")
        pdf_path = os.path.join(tmpdir, "input.pdf")


        # 保存 docx 到临时文件
        with open(docx_path, "wb") as f:
            f.write(docx_bytes)
        

        # 转换 docx 为 pdf，LibreOffice 会自动输出到 tmpdir 中
        try:
            cmd=get_libreoffice_command(tmpdir,docx_path,'pdf')
            result = subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=get_libreoffice_timeout_seconds(), **get_libreoffice_subprocess_kwargs(cmd))
        except subprocess.CalledProcessError as e:
            detail = explain_libreoffice_returncode(e.returncode)
            if detail:
                print(f"LibreOffice诊断: {detail}")
            raise RuntimeError(f"LibreOffice 转换失败: {e.stderr.decode()}")
        except subprocess.TimeoutExpired:
            raise RuntimeError("LibreOffice 转换超时（30秒）")

        # 等待一下，确保文件写入完成
        time.sleep(1)
        

        # 读取转换后的 PDF
        if not os.path.exists(pdf_path):
            raise FileNotFoundError("PDF 文件未生成")

        with open(pdf_path, "rb") as f:
            pdf_content = f.read()
            return pdf_content


def _get_html_from_template(template) -> str:
    """从 TemplateDoc 记录中提取 HTML 模板字符串。"""
    import base64 as _b64
    content_b64 = template.template_content
    if not content_b64:
        return ''
    # DEBUG: 打印原始数据 hex，对比两边是否一致
    print(f"[DEBUG] template={template.name} id={template.id} content_b64 type={type(content_b64)} len={len(content_b64)} hex[:40]={content_b64[:20].hex() if isinstance(content_b64, bytes) else content_b64[:20].encode().hex()}")
    raw_bytes = _b64.b64decode(content_b64)
    print(f"[DEBUG] b64decode len={len(raw_bytes)} hex[:40]={raw_bytes[:20].hex()}")
    for enc in ('utf-8', 'gbk', 'gb2312', 'gb18030', 'latin-1'):
        try:
            return raw_bytes.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    print(f"[_get_html_from_template] 模板 {template.name}(id={template.id}) 无法解码，使用 utf-8 replace 兜底")
    return raw_bytes.decode('utf-8', errors='replace')


def _save_classic_config_file(old_name: str, new_name: str, title_mapping: dict, header_configs: dict, footer_configs: dict, title_list=None, page_num=None, margin=None, orientation=None, line_height=None, py_dir: str = None):
    """
    另存为：从原 .py 配置文件复制，更新 title_mapping / 标题 / 页眉页脚 / 页码。
    header_configs / footer_configs: {'left_header': {...}, 'mid_header': {...}, ...} 六字段配置
    py_dir: 结构目录，默认经典模式 classic_structure_path（向导模式传 structure_path）
    """
    import shutil

    py_dir = py_dir or str(config.classic_structure_path)
    old_path = os.path.join(py_dir, f"{old_name}.py")
    new_path = os.path.join(py_dir, f"{new_name}.py")

    if not os.path.exists(old_path):
        print(f"[_save_classic_config_file] 原配置文件不存在: {old_path}")
        return

    shutil.copy(old_path, new_path)

    with open(new_path, 'r', encoding='utf-8') as f:
        content = f.read()

    import re

    # 1. 重建 title_mapping 段
    lines = ["title_mapping = {\n"]
    keys = list(title_mapping.keys())
    for i, k in enumerate(keys):
        info = title_mapping[k]
        comma = "," if i < len(keys) - 1 else ""
        lines.append(
            f"    '{k}': {{'name': '{info.get('name', '')}', "
            f"'font': '{info.get('font', '宋体')}', "
            f"'font_size': {info.get('font_size', 12)}}}{comma}\n"
        )
    lines.append("}\n")
    content = re.sub(r'title_mapping = \{[\s\S]*?\n\}\n', ''.join(lines), content)

    # 2. 更新 extra_info 中左中右页眉页脚（先删旧行再插入新块）
    _HF_KEYS = ['left_header', 'mid_header', 'right_header', 'left_footer', 'mid_footer', 'right_footer']
    _kept_lines = []
    for _line in content.split('\n'):
        _stripped = _line.strip()
        if any(_stripped.startswith(f"'{k}':") for k in _HF_KEYS):
            continue
        _kept_lines.append(_line)
    content = '\n'.join(_kept_lines)

    hf_lines = []
    for key in _HF_KEYS:
        cfg = (header_configs if 'header' in key else footer_configs).get(key)
        if not cfg or not cfg.get('value'):
            continue
        hf_lines.append(
            f"    '{key}': {{'type': {cfg.get('type', 1)}, 'value': '{cfg.get('value', '')}', "
            f"'font': '{cfg.get('font', '宋体')}', 'font_size': {cfg.get('font_size', 10)}, "
            f"'position': '{cfg.get('position', '')}'}},"
        )
    if hf_lines:
        hf_block = '\n'.join(hf_lines) + '\n'
        # 插入到 extra_info = { 之后
        content = re.sub(r"(extra_info\s*=\s*\{\n)", r"\1" + hf_block, content, count=1)

    # 3. 更新 extra_info 中 title_list
    if title_list:
        tl_lines = ["'title_list': [\n"]
        for t in title_list:
            tl_lines.append(
                f"        {{'name': '报告名称', 'value': '{t.value}', "
                f"'font': '{t.font or '黑体'}', 'font_size': {t.font_size or 15}}},\n"
            )
        tl_lines.append("    ]")
        tl_block = ''.join(tl_lines)
        content = re.sub(r"'title_list'\s*:\s*\[[\s\S]*?\]", tl_block, content)

    # 4. 更新 extra_info 中 page_num（始终写入，position 为空表示不展示）
    if page_num is not None:
        pn_block = f"'page_num': {{'position': '{page_num.get('position', '')}', 'font': '{page_num.get('font', '宋体')}', 'font_size': {page_num.get('font_size', 9)}}}"
        content = re.sub(r"'page_num'\s*:\s*\{[^}]*\}", pn_block, content)

    # 5. 更新 extra_info 中 margin（动态边距）
    if margin is not None:
        content = re.sub(r"'margin'\s*:\s*\{[^}]*\}", f"'margin': {{'value': '{margin}'}}", content)

    # 6. 更新 extra_info 中 orientation（横向/纵向）
    if orientation is not None:
        content = re.sub(r"'orientation'\s*:\s*\{[^}]*\}", f"'orientation': {{'value': {orientation}}}", content)

    # 7. 更新 extra_info 中 line_height（行边距）
    if line_height is not None:
        content = re.sub(r"'line_height'\s*:\s*\{[^}]*\}", f"'line_height': {{'value': '{line_height}'}}", content)

    with open(new_path, 'w', encoding='utf-8') as f:
        f.write(content)

    print(f"[_save_classic_config_file] 配置文件已保存: {new_path}")